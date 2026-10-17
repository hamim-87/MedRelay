"""Tiny slide DSL with two backends: ReportLab (vector PDF) and python-pptx (editable PPTX).

All geometry is in points on a 960x540 page (13.333 x 7.5 in), origin top-left. Text is wrapped with Liberation Sans metrics (metric-compatible with Arial),
so the PDF and the PPTX (font "Arial") break lines the same way. `Deck.warnings` lists every text box whose content does not fit its box.
"""
from __future__ import annotations

import copy
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image, ImageFont

W, H = 960.0, 540.0
FONT_DIR = Path("/usr/share/fonts/truetype/liberation")
FONTS = {("sans", False): FONT_DIR / "LiberationSans-Regular.ttf", ("sans", True): FONT_DIR / "LiberationSans-Bold.ttf",
         ("mono", False): FONT_DIR / "LiberationMono-Regular.ttf", ("mono", True): FONT_DIR / "LiberationMono-Bold.ttf"}
PPTX_FONT = {"sans": "Arial", "mono": "Courier New"}
_FONT_CACHE: dict = {}


def hex2rgb(h: str):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def _pil_font(family: str, bold: bool, size: float):
    key = (family, bold, round(size * 10))
    if key not in _FONT_CACHE:
        _FONT_CACHE[key] = ImageFont.truetype(str(FONTS[(family, bold)]), int(round(size * 10)))
    return _FONT_CACHE[key]


_CHAR_W: dict = {}


def text_width(s: str, family: str, bold: bool, size: float) -> float:
    """Sum of per-glyph advances (no kerning), the same model ReportLab uses when it draws a string."""
    f = _pil_font(family, bold, 100)
    tot = 0.0
    for ch in s:
        k = (family, bold, ch)
        if k not in _CHAR_W:
            _CHAR_W[k] = f.getlength(ch) / 1000.0
        tot += _CHAR_W[k]
    return tot * size


@dataclass
class Run:
    text: str
    bold: bool = False
    color: str | None = None
    size: float | None = None
    family: str | None = None


def _norm_para(p, base_bold, base_color):
    """paragraph -> list[Run]; accepts str | list[str | (str, dict)]"""
    if isinstance(p, str):
        return [Run(p, base_bold, base_color)]
    out = []
    for r in p:
        if isinstance(r, str):
            out.append(Run(r, base_bold, base_color))
        else:
            t, o = r
            out.append(Run(t, o.get("bold", base_bold), o.get("color", base_color), o.get("size"), o.get("family")))
    return out


def wrap_para(runs: list[Run], width: float, size: float, family: str) -> list[list[tuple[Run, str]]]:
    """Greedy word wrap; returns lines, each a list of (run, text-piece)."""
    tokens = []                                   # (run, word, trailing_space)
    for r in runs:
        parts = r.text.split("\n")
        for pi, part in enumerate(parts):
            words = part.split(" ")
            for wi, w in enumerate(words):
                if w == "" and wi > 0 and wi == len(words) - 1:
                    continue
                tokens.append((r, w, wi < len(words) - 1))
            if pi < len(parts) - 1:
                tokens.append((r, "\n", False))
    lines, cur, cur_w = [], [], 0.0
    for r, w, sp in tokens:
        if w == "\n":
            lines.append(cur); cur, cur_w = [], 0.0
            continue
        fam = r.family or family
        sz = r.size or size
        ww = text_width(w, fam, r.bold, sz)
        spw = text_width(" ", fam, r.bold, sz) if sp else 0.0
        if cur and cur_w + ww > width + 0.01:
            lines.append(cur); cur, cur_w = [], 0.0
        cur.append((r, w + (" " if sp else "")))
        cur_w += ww + spw
    lines.append(cur)
    return lines


@dataclass
class Slide:
    ops: list = field(default_factory=list)
    notes: str = ""


class Deck:
    def __init__(self):
        self.slides: list[Slide] = []
        self.warnings: list[str] = []

    def new(self) -> Slide:
        s = Slide(); self.slides.append(s); return s

    # ---- primitives (stored as ops; both backends replay them)
    def bg(self, s, c0="#070D1A", c1="#0E1B38"):
        s.ops.append(("bg", c0, c1))

    def rect(self, s, x, y, w, h, fill=None, line=None, lw=1.0, r=0.0, alpha=1.0, dash=None, line_alpha=1.0):
        s.ops.append(("rect", x, y, w, h, fill, line, lw, r, alpha, dash, line_alpha))

    def ellipse(self, s, x, y, w, h, fill=None, line=None, lw=1.0, alpha=1.0, line_alpha=1.0):
        s.ops.append(("ellipse", x, y, w, h, fill, line, lw, alpha, line_alpha))

    def line(self, s, x1, y1, x2, y2, color="#FFFFFF", lw=1.0, alpha=1.0, dash=None, arrow=False):
        s.ops.append(("line", x1, y1, x2, y2, color, lw, alpha, dash, arrow))

    def image(self, s, path, x, y, w, h, align="center"):
        s.ops.append(("image", str(path), x, y, w, h, align))

    def text(self, s, x, y, w, h, paras, size=14, color="#FFFFFF", bold=False, align="l", lh=1.22, family="sans", anchor="top", space=0.0, name=""):
        if isinstance(paras, str):
            paras = [paras]
        P = [_norm_para(p, bold, color) for p in paras]
        # measure
        total = 0.0
        laid = []
        for pi, runs in enumerate(P):
            lines = wrap_para(runs, w, size, family)
            laid.append(lines)
            for ln in lines:
                mx = max([(r.size or size) for r, _ in ln] + [size])
                total += mx * lh
            if pi < len(P) - 1:
                total += space
        if total > h + 1.5:
            self.warnings.append(f"slide {len(self.slides)}: text '{(name or str(paras[0]))[:40]}' needs {total:.0f}pt but box is {h:.0f}pt")
        s.ops.append(("text", x, y, w, h, P, size, color, bold, align, lh, family, anchor, space, laid, total))
        return total

    # ---- ReportLab backend
    def save_pdf(self, path: str):
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.colors import Color
        pdfmetrics.registerFont(TTFont("LSans", str(FONTS[("sans", False)])))
        pdfmetrics.registerFont(TTFont("LSans-B", str(FONTS[("sans", True)])))
        pdfmetrics.registerFont(TTFont("LMono", str(FONTS[("mono", False)])))
        pdfmetrics.registerFont(TTFont("LMono-B", str(FONTS[("mono", True)])))
        fmap = {("sans", False): "LSans", ("sans", True): "LSans-B", ("mono", False): "LMono", ("mono", True): "LMono-B"}

        def col(h, a=1.0):
            r, g, b = hex2rgb(h)
            return Color(r / 255, g / 255, b / 255, alpha=a)
        c = canvas.Canvas(path, pagesize=(W, H))
        c.setTitle("MedRelay - IBM Z Datathon 2026 pitch deck")
        c.setAuthor("MedRelay team")
        for s in self.slides:
            for op in s.ops:
                k = op[0]
                if k == "bg":
                    _, c0, c1 = op
                    c.saveState(); p = c.beginPath(); p.rect(0, 0, W, H); c.clipPath(p, stroke=0, fill=0)
                    c.linearGradient(0, H, W, 0, (col(c0), col(c1)), (0, 1)); c.restoreState()
                elif k == "rect":
                    _, x, y, w, h, fill, line, lw, r, alpha, dash, la = op
                    if fill: c.setFillColor(col(fill, alpha))
                    if line: c.setStrokeColor(col(line, la)); c.setLineWidth(lw)
                    if dash: c.setDash(*dash)
                    if r: c.roundRect(x, H - y - h, w, h, r, stroke=1 if line else 0, fill=1 if fill else 0)
                    else: c.rect(x, H - y - h, w, h, stroke=1 if line else 0, fill=1 if fill else 0)
                    c.setDash()
                elif k == "ellipse":
                    _, x, y, w, h, fill, line, lw, alpha, la = op
                    if fill: c.setFillColor(col(fill, alpha))
                    if line: c.setStrokeColor(col(line, la)); c.setLineWidth(lw)
                    c.ellipse(x, H - y - h, x + w, H - y, stroke=1 if line else 0, fill=1 if fill else 0)
                elif k == "line":
                    _, x1, y1, x2, y2, color, lw, alpha, dash, arrow = op
                    c.setStrokeColor(col(color, alpha)); c.setLineWidth(lw)
                    if dash: c.setDash(*dash)
                    import math
                    ang = math.atan2(y2 - y1, x2 - x1)
                    ex, ey = x2, y2
                    if arrow:
                        ex, ey = x2 - 6 * math.cos(ang), y2 - 6 * math.sin(ang)
                    c.line(x1, H - y1, ex, H - ey); c.setDash()
                    if arrow:
                        c.setFillColor(col(color, alpha)); p = c.beginPath()
                        hl, hw = 8.0, 3.6
                        p.moveTo(x2, H - y2)
                        p.lineTo(x2 - hl * math.cos(ang) + hw * math.sin(ang), H - (y2 - hl * math.sin(ang) - hw * math.cos(ang)))
                        p.lineTo(x2 - hl * math.cos(ang) - hw * math.sin(ang), H - (y2 - hl * math.sin(ang) + hw * math.cos(ang)))
                        p.close(); c.drawPath(p, stroke=0, fill=1)
                elif k == "image":
                    _, path_, x, y, w, h, align = op
                    im = Image.open(path_); iw, ih = im.size
                    sc = min(w / iw, h / ih); dw, dh = iw * sc, ih * sc
                    dx = x + (w - dw) / 2 if align == "center" else x
                    dy = y + (h - dh) / 2 if align == "center" else y
                    c.drawImage(path_, dx, H - dy - dh, dw, dh, mask="auto")
                elif k == "text":
                    _, x, y, w, h, P, size, color, bold, align, lh, family, anchor, space, laid, total = op
                    cy = y
                    if anchor == "middle": cy = y + (h - total) / 2
                    elif anchor == "bottom": cy = y + h - total
                    for pi, lines in enumerate(laid):
                        for ln in lines:
                            mx = max([(r.size or size) for r, _ in ln] + [size])
                            base = cy + mx * lh * 0.5 + mx * 0.32          # visual baseline inside the line box
                            wline = sum(text_width(t, r.family or family, r.bold, r.size or size) for r, t in ln)
                            xx = x if align == "l" else (x + (w - wline) / 2 if align == "c" else x + w - wline)
                            for r, t in ln:
                                fam, sz = r.family or family, r.size or size
                                c.setFont(fmap[(fam, r.bold)], sz); c.setFillColor(col(r.color or color))
                                c.drawString(xx, H - base, t.rstrip("\n"))
                                xx += text_width(t, fam, r.bold, sz)
                            cy += mx * lh
                        if pi < len(laid) - 1: cy += space
            c.showPage()
        c.save()

    # ---- python-pptx backend
    def save_pptx(self, path: str):
        from pptx import Presentation
        from pptx.util import Emu, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.oxml.ns import qn
        from lxml import etree
        E = lambda v: Emu(int(round(v * 12700)))
        prs = Presentation(); prs.slide_width = E(W); prs.slide_height = E(H)
        blank = prs.slide_layouts[6]

        def rgb(h): return RGBColor(*hex2rgb(h))

        def set_alpha(fill_fmt, alpha):
            if alpha >= 0.999: return
            srgb = fill_fmt._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = etree.SubElement(srgb, qn("a:alpha")); a.set("val", str(int(alpha * 100000)))

        def set_line_alpha(shape, alpha):
            if alpha >= 0.999: return
            ln = shape.line._get_or_add_ln()
            srgb = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = etree.SubElement(srgb, qn("a:alpha")); a.set("val", str(int(alpha * 100000)))

        dash_map = {(4, 3): MSO_LINE_DASH_STYLE.DASH, (2, 3): MSO_LINE_DASH_STYLE.ROUND_DOT, (6, 4): MSO_LINE_DASH_STYLE.DASH}
        for s in self.slides:
            sl = prs.slides.add_slide(blank)
            if s.notes:
                sl.notes_slide.notes_text_frame.text = s.notes
            for op in s.ops:
                k = op[0]
                if k == "bg":
                    _, c0, c1 = op
                    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, E(W), E(H)); sh.line.fill.background()
                    sh.fill.gradient(); sh.fill.gradient_angle = 330
                    st = sh.fill.gradient_stops; st[0].color.rgb = rgb(c0); st[0].position = 0; st[1].color.rgb = rgb(c1); st[1].position = 1
                    sh.name = "background"
                elif k == "rect":
                    _, x, y, w, h, fill, line, lw, r, alpha, dash, la = op
                    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
                    if r: sh.adjustments[0] = min(0.5, r / min(w, h))
                    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); set_alpha(sh.fill, alpha)
                    else: sh.fill.background()
                    if line:
                        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw); set_line_alpha(sh, la)
                        if dash: sh.line.dash_style = dash_map.get(tuple(dash), MSO_LINE_DASH_STYLE.DASH)
                    else: sh.line.fill.background()
                    sh.shadow.inherit = False
                elif k == "ellipse":
                    _, x, y, w, h, fill, line, lw, alpha, la = op
                    sh = sl.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(y), E(w), E(h))
                    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); set_alpha(sh.fill, alpha)
                    else: sh.fill.background()
                    if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw); set_line_alpha(sh, la)
                    else: sh.line.fill.background()
                    sh.shadow.inherit = False
                elif k == "line":
                    _, x1, y1, x2, y2, color, lw, alpha, dash, arrow = op
                    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
                    cn.line.color.rgb = rgb(color); cn.line.width = Pt(lw); set_line_alpha(cn, alpha)
                    if dash: cn.line.dash_style = dash_map.get(tuple(dash), MSO_LINE_DASH_STYLE.DASH)
                    if arrow:
                        ln = cn.line._get_or_add_ln(); te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
                elif k == "image":
                    _, path_, x, y, w, h, align = op
                    im = Image.open(path_); iw, ih = im.size
                    sc = min(w / iw, h / ih); dw, dh = iw * sc, ih * sc
                    dx = x + (w - dw) / 2 if align == "center" else x
                    dy = y + (h - dh) / 2 if align == "center" else y
                    sl.shapes.add_picture(path_, E(dx), E(dy), E(dw), E(dh))
                elif k == "text":
                    _, x, y, w, h, P, size, color, bold, align, lh, family, anchor, space, laid, total = op
                    tb = sl.shapes.add_textbox(E(x), E(y), E(w + 2), E(h))
                    tf = tb.text_frame; tf.word_wrap = True
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
                    for pi, runs in enumerate(P):
                        para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
                        mx = max([(r.size or size) for r in runs] + [size])
                        para.line_spacing = Pt(mx * lh)
                        if pi < len(P) - 1 and space: para.space_after = Pt(space)
                        for r in runs:
                            parts = r.text.split("\n")
                            for qi, part in enumerate(parts):
                                run = para.add_run(); run.text = part
                                f = run.font; f.name = PPTX_FONT[r.family or family]; f.size = Pt(r.size or size); f.bold = bool(r.bold); f.color.rgb = rgb(r.color or color)
                                if qi < len(parts) - 1:
                                    para.add_line_break()
        prs.save(path)

    def check_pptx(self, path: str) -> list[str]:
        """Programmatic geometry check of the saved PPTX (it cannot be rendered here): every shape must lie inside the slide."""
        from pptx import Presentation
        prs = Presentation(path); issues = []
        for i, sl in enumerate(prs.slides, 1):
            for sh in sl.shapes:
                if sh.left is None: continue
                if sh.left < -1000 or sh.top < -1000 or sh.left + sh.width > prs.slide_width + 30000 or sh.top + sh.height > prs.slide_height + 30000:
                    issues.append(f"slide {i}: shape '{sh.name}' leaves the slide")
        return issues
